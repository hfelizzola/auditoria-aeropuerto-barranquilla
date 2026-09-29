"""
lote_excel.py
=============
Lectura y escritura del Excel de lote de validación, respetando EXACTAMENTE el
layout que se entrega manualmente:

- Hoja "Resumen": bloque de título/notas (filas 1-5, texto libre), encabezado
  (normalmente fila 6, desde la columna B) con las 11 columnas
  `#`, `Fila UNIVERSALIDAD`, `OP`, `Tercero (según UNIVERSALIDAD)`, `Cuenta contable`,
  `Fecha`, `Valor (COP)`, `Moneda carpeta`, `URL soporte`, `Resultado de validación`,
  `Observación`, y una fila por transacción.
- Hoja "Hallazgos detallados": tabla de 2 columnas (título / detalle).
- Hoja "Control de calidad": anomalías de los datos de UNIVERSALIDAD.

El código de auditoría solo escribe en `Resultado de validación` y `Observación`
de filas existentes; el resto de columnas no se toca. Cada guardado es atómico
(archivo temporal + reemplazo) para que una interrupción nunca deje el .xlsx corrupto.
"""

from __future__ import annotations

import logging
import os
import re
import time
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

HOJA_RESUMEN = "Resumen"
HOJA_HALLAZGOS = "Hallazgos detallados"
HOJA_CONTROL = "Control de calidad"

# clave interna -> encabezado exacto en la hoja Resumen (en orden)
COLUMNAS = [
    ("rank", "#"),
    ("fila_universalidad", "Fila UNIVERSALIDAD"),
    ("op", "OP"),
    ("tercero", "Tercero (según UNIVERSALIDAD)"),
    ("cuenta", "Cuenta contable"),
    ("fecha", "Fecha"),
    ("valor", "Valor (COP)"),
    ("moneda", "Moneda carpeta"),
    ("url", "URL soporte"),
    ("resultado", "Resultado de validación"),
    ("observacion", "Observación"),
]
ENCABEZADO_POR_CLAVE = dict(COLUMNAS)
COLUMNA_INICIAL = 2  # la tabla empieza en la columna B
FILA_ENCABEZADO_DEFECTO = 6
ENCABEZADO_HALLAZGOS = ("Hallazgos (rank / OP)", "Detalle")
MARCADOR_CONTROL_AUTO = "Anomalías detectadas automáticamente por el orquestador"
FORMATO_VALOR = "$#,##0;($#,##0);-"
MAX_CARACTERES_CELDA = 32000

AZUL = "1F4E79"
GRIS = "595959"
ANCHOS_RESUMEN = {"A": 2, "B": 6, "C": 14, "D": 10, "E": 38, "F": 22, "G": 12, "H": 18,
                  "I": 10, "J": 60, "K": 22, "L": 90}


def limpiar_texto_celda(valor: Any) -> Any:
    if isinstance(valor, str):
        valor = ILLEGAL_CHARACTERS_RE.sub("", valor)
        if len(valor) > MAX_CARACTERES_CELDA:
            valor = valor[: MAX_CARACTERES_CELDA - 20] + " …[truncado]"
    return valor


@dataclass
class FilaLote:
    fila_excel: int
    rank: Optional[int]
    fila_universalidad: Optional[int]
    op: Any
    tercero: Any
    cuenta: Any
    fecha: Any
    valor: Any
    moneda: Any
    url: Optional[str]
    resultado: Optional[str]
    observacion: Optional[str]

    @property
    def validada(self) -> bool:
        return bool(self.resultado and str(self.resultado).strip())


def _a_int(v: Any) -> Optional[int]:
    try:
        if v is None or str(v).strip() == "":
            return None
        return int(float(str(v).strip().lstrip("#")))
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------------------
# Hipervínculos embebidos (respaldo cuando la celda no trae la URL como texto)
# ---------------------------------------------------------------------------

_NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}


def _resolver_ruta_zip(base: str, destino: str) -> str:
    if destino.startswith("/"):
        return destino.lstrip("/")
    partes: List[str] = []
    for p in (PurePosixPath(base).parent / destino).parts:
        if p == "..":
            if partes:
                partes.pop()
        elif p != ".":
            partes.append(p)
    return "/".join(partes)


def extraer_hipervinculos_xlsx(ruta: Path, hoja: str) -> Dict[str, str]:
    """{referencia de celda ('J7') -> URL destino} leyendo el paquete OOXML directamente.

    openpyxl en modo read_only/data_only solo devuelve el texto visible de la celda,
    no el destino del hipervínculo. Aquí se recorre:
    xl/workbook.xml -> xl/_rels/workbook.xml.rels -> XML de la hoja -> _rels/<hoja>.xml.rels
    y se mapea el r:id de cada <hyperlink> a su Target.
    """
    resultado: Dict[str, str] = {}
    with zipfile.ZipFile(ruta) as zf:
        wb_xml = ET.fromstring(zf.read("xl/workbook.xml"))
        rid_hoja = None
        for s in wb_xml.iterfind("m:sheets/m:sheet", _NS):
            if s.get("name") == hoja:
                rid_hoja = s.get(f"{{{_NS['r']}}}id")
                break
        if rid_hoja is None:
            raise ValueError(f"La hoja '{hoja}' no existe en '{ruta.name}'.")
        rels_wb = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        destino_hoja = next((r.get("Target") for r in rels_wb.iterfind("rel:Relationship", _NS)
                             if r.get("Id") == rid_hoja), None)
        if destino_hoja is None:
            return resultado
        ruta_hoja = _resolver_ruta_zip("xl/workbook.xml", destino_hoja)
        ruta_rels_hoja = str(PurePosixPath(ruta_hoja).parent / "_rels" / (PurePosixPath(ruta_hoja).name + ".rels"))
        if ruta_rels_hoja not in zf.namelist():
            return resultado
        rels_hoja = {r.get("Id"): r.get("Target")
                     for r in ET.fromstring(zf.read(ruta_rels_hoja)).iterfind("rel:Relationship", _NS)
                     if r.get("TargetMode") == "External" or str(r.get("Type", "")).endswith("/hyperlink")}
        # iterparse: las hojas grandes (UNIVERSALIDAD) no se cargan completas en memoria
        with zf.open(ruta_hoja) as fh:
            for _, elem in ET.iterparse(fh, events=("end",)):
                if elem.tag == f"{{{_NS['m']}}}hyperlink":
                    rid = elem.get(f"{{{_NS['r']}}}id")
                    ref = elem.get("ref")
                    if rid in rels_hoja and ref:
                        # un hipervínculo puede cubrir un rango (A1:A3): se asigna a la primera celda
                        resultado[ref.split(":")[0]] = rels_hoja[rid]
                elif elem.tag == f"{{{_NS['m']}}}row":
                    elem.clear()
    return resultado


# ---------------------------------------------------------------------------
# Libro de lote
# ---------------------------------------------------------------------------

class LoteExcel:
    def __init__(self, ruta: Path, wb: openpyxl.Workbook):
        self.ruta = Path(ruta)
        self.wb = wb
        self._ws = wb[HOJA_RESUMEN]
        self.fila_encabezado, self.col = self._mapear_encabezado()
        self._hipervinculos: Optional[Dict[str, str]] = None

    # ----- creación / apertura -------------------------------------------

    @classmethod
    def abrir(cls, ruta: Path) -> "LoteExcel":
        ruta = Path(ruta)
        wb = openpyxl.load_workbook(ruta)  # sin data_only: conserva fórmulas al guardar
        if HOJA_RESUMEN not in wb.sheetnames:
            raise ValueError(f"'{ruta.name}' no tiene la hoja '{HOJA_RESUMEN}'. Hojas: {wb.sheetnames}")
        lote = cls(ruta, wb)
        lote._asegurar_hojas_auxiliares()
        return lote

    @classmethod
    def crear(cls, ruta: Path, titulo: str, notas: Sequence[str]) -> "LoteExcel":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = HOJA_RESUMEN
        for letra, ancho in ANCHOS_RESUMEN.items():
            ws.column_dimensions[letra].width = ancho
        ultima_letra_notas = get_column_letter(COLUMNA_INICIAL + len(COLUMNAS) - 3)  # B:J
        ws.cell(row=2, column=COLUMNA_INICIAL, value=limpiar_texto_celda(titulo)).font = Font(bold=True, size=13, color=AZUL)
        ws.merge_cells(f"B2:{ultima_letra_notas}2")
        for i, nota in enumerate(list(notas)[:3]):
            fila = 3 + i
            c = ws.cell(row=fila, column=COLUMNA_INICIAL, value=limpiar_texto_celda(nota))
            c.font = Font(color=GRIS)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(f"B{fila}:{ultima_letra_notas}{fila}")
            ws.row_dimensions[fila].height = max(30, min(150, 15 * (len(nota) // 160 + 1)))
        for j, (_, etiqueta) in enumerate(COLUMNAS):
            c = ws.cell(row=FILA_ENCABEZADO_DEFECTO, column=COLUMNA_INICIAL + j, value=etiqueta)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=AZUL)
            c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = f"B{FILA_ENCABEZADO_DEFECTO + 1}"
        lote = cls(ruta, wb)
        lote._asegurar_hojas_auxiliares()
        return lote

    def _asegurar_hojas_auxiliares(self) -> None:
        if HOJA_HALLAZGOS not in self.wb.sheetnames:
            ws = self.wb.create_sheet(HOJA_HALLAZGOS)
            ws.column_dimensions["A"].width = 2
            ws.column_dimensions["B"].width = 45
            ws.column_dimensions["C"].width = 110
            ws["B2"] = "Hallazgos que requieren atención"
            ws["B2"].font = Font(bold=True, size=13, color=AZUL)
            ws["B3"] = ("Entradas agrupadas por patrón. Las marcadas [Automático] las genera y actualiza "
                        "el orquestador al cerrar cada sub-lote; el resto es texto libre del auditor.")
            ws["B3"].font = Font(color=GRIS)
            for j, txt in enumerate(ENCABEZADO_HALLAZGOS):
                c = ws.cell(row=5, column=2 + j, value=txt)
                c.font = Font(bold=True, color=AZUL)
        if HOJA_CONTROL not in self.wb.sheetnames:
            ws = self.wb.create_sheet(HOJA_CONTROL)
            for letra, ancho in {"A": 2, "B": 10, "C": 10, "D": 38, "E": 30, "F": 90, "G": 14}.items():
                ws.column_dimensions[letra].width = ancho
            ws["B2"] = "Control de calidad — anomalías de los datos de UNIVERSALIDAD (documentadas, no corregidas aquí)"
            ws["B2"].font = Font(bold=True, size=13)

    def _mapear_encabezado(self) -> Tuple[int, Dict[str, int]]:
        ws = self._ws
        for r in range(1, min(ws.max_row, 40) + 1):
            valores = {str(ws.cell(row=r, column=c).value).strip(): c
                       for c in range(1, ws.max_column + 1) if ws.cell(row=r, column=c).value is not None}
            if "#" in valores and ENCABEZADO_POR_CLAVE["url"] in valores:
                col = {clave: valores[etq] for clave, etq in COLUMNAS if etq in valores}
                faltan = {"rank", "op", "url", "resultado", "observacion"} - set(col)
                if faltan:
                    raise ValueError(f"Encabezado incompleto en '{self.ruta.name}': faltan {sorted(faltan)}")
                return r, col
        raise ValueError(f"No se encontró la fila de encabezado ('#' y 'URL soporte') en '{self.ruta.name}'.")

    # ----- lectura -------------------------------------------------------

    def _url_de_celda(self, fila_excel: int) -> Optional[str]:
        celda = self._ws.cell(row=fila_excel, column=self.col["url"])
        valor = str(celda.value).strip() if celda.value is not None else ""
        if valor.lower().startswith(("http://", "https://")):
            return valor
        if celda.hyperlink is not None and celda.hyperlink.target:
            return celda.hyperlink.target
        if self.ruta.exists():  # respaldo: hipervínculo embebido leído del paquete OOXML
            if self._hipervinculos is None:
                try:
                    self._hipervinculos = extraer_hipervinculos_xlsx(self.ruta, HOJA_RESUMEN)
                except Exception as e:  # noqa: BLE001 - el respaldo nunca debe romper la lectura
                    logger.debug(f"No se pudieron leer hipervínculos embebidos: {e}")
                    self._hipervinculos = {}
            return self._hipervinculos.get(celda.coordinate)
        return None

    def filas(self) -> List[FilaLote]:
        ws = self._ws
        salida: List[FilaLote] = []
        for r in range(self.fila_encabezado + 1, ws.max_row + 1):
            v = {clave: ws.cell(row=r, column=c).value for clave, c in self.col.items()}
            if v.get("rank") is None and v.get("op") is None:
                continue  # fila vacía / separador
            salida.append(FilaLote(
                fila_excel=r, rank=_a_int(v.get("rank")), fila_universalidad=_a_int(v.get("fila_universalidad")),
                op=v.get("op"), tercero=v.get("tercero"), cuenta=v.get("cuenta"), fecha=v.get("fecha"),
                valor=v.get("valor"), moneda=v.get("moneda"), url=self._url_de_celda(r),
                resultado=v.get("resultado"), observacion=v.get("observacion"),
            ))
        return salida

    def ranks_presentes(self) -> set:
        return {f.rank for f in self.filas() if f.rank is not None}

    # ----- escritura -----------------------------------------------------

    def agregar_filas(self, registros: Iterable[Dict[str, Any]]) -> int:
        """Agrega filas nuevas al final de la tabla (solo columnas de datos)."""
        ws = self._ws
        ultima = max([f.fila_excel for f in self.filas()] or [self.fila_encabezado])
        n = 0
        for reg in registros:
            ultima += 1
            n += 1
            for clave, c in self.col.items():
                if clave in ("resultado", "observacion"):
                    celda = ws.cell(row=ultima, column=c)
                    celda.alignment = Alignment(wrap_text=True, vertical="top")
                    continue
                celda = ws.cell(row=ultima, column=c, value=limpiar_texto_celda(reg.get(clave)))
                celda.alignment = Alignment(wrap_text=True, vertical="top")
                if clave == "valor":
                    celda.number_format = FORMATO_VALOR
        return n

    def escribir_resultado(self, fila_excel: int, etiqueta: str, observacion: str) -> None:
        self._ws.cell(row=fila_excel, column=self.col["resultado"], value=etiqueta)
        self._ws.cell(row=fila_excel, column=self.col["observacion"], value=limpiar_texto_celda(observacion))

    def actualizar_hallazgos(self, entradas: Sequence[Tuple[str, str]]) -> None:
        """Actualiza (por título exacto) o agrega entradas en 'Hallazgos detallados'."""
        ws = self.wb[HOJA_HALLAZGOS]
        fila_enc = None
        for r in range(1, ws.max_row + 1):
            if str(ws.cell(row=r, column=2).value or "").strip() == ENCABEZADO_HALLAZGOS[0]:
                fila_enc = r
        existentes = {str(ws.cell(row=r, column=2).value).strip(): r
                      for r in range((fila_enc or 0) + 1, ws.max_row + 1) if ws.cell(row=r, column=2).value}
        if fila_enc is None:
            fila_enc = ws.max_row + 2
            for j, txt in enumerate(ENCABEZADO_HALLAZGOS):
                ws.cell(row=fila_enc, column=2 + j, value=txt).font = Font(bold=True, color=AZUL)
        siguiente = max([fila_enc] + list(existentes.values())) + 1
        for titulo, detalle in entradas:
            r = existentes.get(titulo)
            if r is None:
                r = siguiente
                siguiente += 1
            ws.cell(row=r, column=2, value=limpiar_texto_celda(titulo)).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=r, column=3, value=limpiar_texto_celda(detalle)).alignment = Alignment(wrap_text=True, vertical="top")

    def reescribir_control_automatico(self, encabezados: Sequence[str], filas: Sequence[Sequence[Any]],
                                      subtitulo: str = "") -> None:
        """Reemplaza la sección automática (siempre al final) de 'Control de calidad'.

        Las secciones manuales que estén ANTES del marcador no se tocan.
        """
        ws = self.wb[HOJA_CONTROL]
        inicio = None
        for r in range(1, ws.max_row + 1):
            if str(ws.cell(row=r, column=2).value or "").startswith(MARCADOR_CONTROL_AUTO):
                inicio = r
                break
        if inicio is not None:
            ws.delete_rows(inicio, ws.max_row - inicio + 1)
        if not filas:
            return
        inicio = ws.max_row + 2
        ws.cell(row=inicio, column=2, value=f"{MARCADOR_CONTROL_AUTO} (no corregidas aquí)").font = Font(bold=True)
        if subtitulo:
            ws.cell(row=inicio + 1, column=2, value=limpiar_texto_celda(subtitulo)).font = Font(color=GRIS)
        fila_enc = inicio + 2
        for j, txt in enumerate(encabezados):
            ws.cell(row=fila_enc, column=2 + j, value=txt).font = Font(bold=True)
        for i, fila in enumerate(filas, start=1):
            for j, v in enumerate(fila):
                c = ws.cell(row=fila_enc + i, column=2 + j, value=limpiar_texto_celda(v))
                c.alignment = Alignment(wrap_text=True, vertical="top")

    def guardar(self, reintentos: int = 6) -> None:
        """Guardado atómico: escribe a un temporal y lo reemplaza (resiste interrupciones)."""
        self.ruta.parent.mkdir(parents=True, exist_ok=True)
        tmp = self.ruta.with_name(f".~{self.ruta.stem}.tmp.xlsx")
        for intento in range(1, reintentos + 1):
            try:
                self.wb.save(tmp)
                os.replace(tmp, self.ruta)
                return
            except PermissionError:
                if intento == reintentos:
                    raise PermissionError(
                        f"No se pudo guardar '{self.ruta.name}': ¿está abierto en Excel? Ciérrelo y reanude.")
                logger.warning(f"'{self.ruta.name}' bloqueado (¿abierto en Excel?); reintento en 5s...")
                time.sleep(5)


def buscar_lote_mas_reciente(directorio: Path) -> Optional[Path]:
    candidatos = [p for p in Path(directorio).glob("Validacion_Soportes_*.xlsx") if not p.name.startswith((".~", "~$"))]
    return max(candidatos, key=lambda p: p.stat().st_mtime) if candidatos else None


def slug_lote(ruta: Path) -> str:
    return re.sub(r"[^\w.-]+", "_", Path(ruta).stem)
