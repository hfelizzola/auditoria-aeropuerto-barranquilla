"""Layout del Excel de lote: creación, lectura, escritura solo en resultado/observación,
hallazgos por título, sección automática de control de calidad e hipervínculos embebidos."""

import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from auditoria.lote_excel import (  # noqa: E402
    COLUMNAS, HOJA_CONTROL, HOJA_HALLAZGOS, LoteExcel, MARCADOR_CONTROL_AUTO, extraer_hipervinculos_xlsx,
)

FILA = {"rank": 3505, "fila_universalidad": 3708, "op": 11745, "tercero": "B&S INGENIERIA SAS",
        "cuenta": "MANTENIMIENTO", "fecha": "2024-03-04", "valor": -12978293, "moneda": "COP",
        "url": "https://aerobaq.sharepoint.com/sites/FTAPA/x.pdf"}


class LayoutLote(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.ruta = Path(self.tmp.name) / "Validacion_Soportes_Rango1-2_UNIVERSALIDAD_ABAS1.xlsx"

    def tearDown(self):
        self.tmp.cleanup()

    def _crear(self):
        lote = LoteExcel.crear(self.ruta, "Validación de soportes — prueba", ["nota 1", "nota 2"])
        lote.agregar_filas([FILA, dict(FILA, rank=3506, op=10192, url=None)])
        lote.guardar()
        return LoteExcel.abrir(self.ruta)

    def test_estructura_exacta(self):
        lote = self._crear()
        ws = lote.wb["Resumen"]
        self.assertEqual(lote.fila_encabezado, 6)
        self.assertEqual([ws.cell(row=6, column=2 + j).value for j in range(len(COLUMNAS))], [e for _, e in COLUMNAS])
        self.assertEqual(lote.wb.sheetnames, ["Resumen", HOJA_HALLAZGOS, HOJA_CONTROL])
        filas = lote.filas()
        self.assertEqual([f.rank for f in filas], [3505, 3506])
        self.assertEqual(filas[0].valor, -12978293)
        self.assertFalse(filas[0].validada)

    def test_escribe_solo_resultado_y_observacion(self):
        lote = self._crear()
        antes = [[c.value for c in fila] for fila in lote.wb["Resumen"].iter_rows(min_row=7, max_col=10)]
        f = lote.filas()[0]
        lote.escribir_resultado(f.fila_excel, "COHERENTE", "Obs.")
        lote.guardar()
        lote = LoteExcel.abrir(self.ruta)
        despues = [[c.value for c in fila] for fila in lote.wb["Resumen"].iter_rows(min_row=7, max_col=10)]
        self.assertEqual(antes, despues)
        self.assertEqual((lote.filas()[0].resultado, lote.filas()[0].observacion), ("COHERENTE", "Obs."))
        self.assertTrue(lote.filas()[0].validada)
        self.assertEqual(lote.ranks_presentes(), {3505, 3506})

    def test_hallazgos_se_actualizan_por_titulo(self):
        lote = self._crear()
        lote.actualizar_hallazgos([("[Automático] RETEG", "3 casos"), ("[Automático] USD", "1 caso")])
        lote.actualizar_hallazgos([("[Automático] RETEG", "4 casos")])
        ws = lote.wb[HOJA_HALLAZGOS]
        valores = [(ws.cell(row=r, column=2).value, ws.cell(row=r, column=3).value) for r in range(1, ws.max_row + 1)]
        self.assertEqual(sum(1 for t, _ in valores if t == "[Automático] RETEG"), 1)
        self.assertIn(("[Automático] RETEG", "4 casos"), valores)

    def test_control_automatico_reemplaza_su_seccion(self):
        lote = self._crear()
        ws = lote.wb[HOJA_CONTROL]
        ws["B5"] = "Sección manual"
        lote.reescribir_control_automatico(["Rank", "OP"], [[1, "10"], [2, "20"]])
        lote.reescribir_control_automatico(["Rank", "OP"], [[3, "30"]])
        textos = [ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)]
        self.assertEqual(sum(1 for t in textos if str(t or "").startswith(MARCADOR_CONTROL_AUTO)), 1)
        self.assertIn("Sección manual", textos)
        self.assertIn(3, textos)
        self.assertNotIn(1, textos)

    def test_hipervinculo_embebido(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "UNIVERSALIDAD"
        ws["A1"] = "LINK"
        ws["A2"] = "Ver soporte"
        ws["A2"].hyperlink = "https://aerobaq.sharepoint.com/:b:/s/FTAPA/abc"
        ruta = Path(self.tmp.name) / "h.xlsx"
        wb.save(ruta)
        self.assertEqual(extraer_hipervinculos_xlsx(ruta, "UNIVERSALIDAD"),
                         {"A2": "https://aerobaq.sharepoint.com/:b:/s/FTAPA/abc"})

    def test_lote_manual_real_si_existe(self):
        real = Path(__file__).resolve().parent.parent / "Validacion_Soportes_Rango3501-4000_UNIVERSALIDAD_ABAS1_RESULTADO.xlsx"
        if not real.exists():
            self.skipTest("lote manual no disponible (no se versiona)")
        lote = LoteExcel.abrir(real)
        filas = lote.filas()
        self.assertEqual(len(filas), 501)
        self.assertTrue(all(f.url and f.url.startswith("https://") for f in filas))


if __name__ == "__main__":
    unittest.main()
